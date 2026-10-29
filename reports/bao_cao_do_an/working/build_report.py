from pathlib import Path
import re,json,copy,shutil,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(r'D:\chatbot-law');OUT=ROOT/'reports/bao_cao_do_an';WORK=OUT/'working'
REF=Path(r'C:\Users\Asus\Desktop\Bao_cao_Do_an_co_so_POS_Quan_ly_kho_Canh_bao_ton_kho_thap_FINAL.docx')
FINAL=OUT/'Bao_cao_Hoc_phan_Ung_dung_AI_Mai_Tran_Thien_Tam.docx'
audit=json.loads((WORK/'audit.json').read_text(encoding='utf-8'))
shutil.copy2(REF,FINAL)
d=Document(FINAL)
body=d._element.body
for e in list(body)[16:]:
 if e.tag!=qn('w:sectPr'):body.remove(e)

def font(run,size=14,bold=False,italic=False):
 run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold;run.font.italic=italic;run.font.color.rgb=RGBColor(0,0,0)
 rpr=run._r.get_or_add_rPr();rfonts=rpr.find(qn('w:rFonts'))
 if rfonts is not None:
  for k in ['ascii','hAnsi','eastAsia','cs']:rfonts.set(qn('w:'+k),'Times New Roman')

def set_style(name,size,bold=False):
 s=d.styles[name] if name in d.styles else d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
 s._element.get_or_add_rPr().append(OxmlElement('w:rFonts')) if s._element.get_or_add_rPr().find(qn('w:rFonts')) is None else None
 rf=s._element.get_or_add_rPr().find(qn('w:rFonts'))
 for k in list(rf.attrib):
  if 'Theme' in k or 'theme' in k:del rf.attrib[k]
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Times New Roman')
 return s

s=set_style('Normal',14);s.paragraph_format.line_spacing=1.5;s.paragraph_format.space_after=Pt(6);s.paragraph_format.space_before=Pt(0);s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;s.paragraph_format.widow_control=True
for name,size in [('Heading 1',16),('Heading 2',14),('Heading 3',14)]:
 s=set_style(name,size,True);s.paragraph_format.line_spacing=1.5;s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=True
d.styles['Heading 1'].paragraph_format.page_break_before=True;d.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name,size in [('Front Heading',16),('Hinh',14),('Bang',14),('Code',10),('Table Text',14),('Title',25),('Field Container',1)]:set_style(name,size,name in ['Front Heading','Bang','Title'])
d.styles['Field Container'].paragraph_format.line_spacing=Pt(1);d.styles['Field Container'].paragraph_format.space_after=Pt(0)
d.styles['Front Heading'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
d.styles['Front Heading'].paragraph_format.space_after=Pt(14)
for name in ['Hinh','Bang']:
 d.styles[name].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;d.styles[name].paragraph_format.line_spacing=1.1;d.styles[name].paragraph_format.space_after=Pt(8)
d.styles['Bang'].paragraph_format.keep_with_next=True
d.styles['Hinh'].font.italic=True
d.styles['Table Text'].paragraph_format.line_spacing=1.12;d.styles['Table Text'].paragraph_format.space_after=Pt(3);d.styles['Table Text'].paragraph_format.space_before=Pt(3)
d.styles['Code'].paragraph_format.line_spacing=1.1;d.styles['Code'].paragraph_format.space_after=Pt(3);d.styles['Code'].font.name='Consolas'
for i in range(1,4):
 s=set_style('TOC '+str(i),14);s.paragraph_format.line_spacing=1.1;s.paragraph_format.space_after=Pt(4)

topic='XÂY DỰNG CHATBOT HỖ TRỢ TRA CỨU QUY ĐỊNH PHÁP LUẬT LAO ĐỘNG VIỆT NAM'
for p in d.paragraphs:
 if p.text=='BÁO CÁO ĐỒ ÁN CƠ SỞ':
  p.text='BÁO CÁO HỌC PHẦN\nỨNG DỤNG AI'
  for r in p.runs:font(r,16,True)
  p.paragraph_format.space_before=Pt(90);p.paragraph_format.space_after=Pt(15);p.paragraph_format.line_spacing=1.15
 elif p.text.startswith('POS BÁN HÀNG'):
  p.text=topic;p.style=d.styles['Title'];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for r in p.runs:font(r,25,True)
  p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(36);p.paragraph_format.line_spacing=1.1
for t in d.tables:
 vals=['Người hướng dẫn: [CẦN ĐIỀN]','Người thực hiện: MAI TRẦN THIỆN TÂM','Mã số sinh viên: 2305CT2070','Lớp: [CẦN ĐIỀN]','Khóa: [CẦN ĐIỀN]']
 while len(t.rows)<5:t.add_row()
 for row,txt in zip(t.rows,vals):
  row.cells[0].text=txt
  p=row.cells[0].paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(5)
  for r in p.runs:font(r,14,True)
 for row in t.rows:row.height=None

def configure_section(s,body_style=True):
 s.page_width=Cm(21);s.page_height=Cm(29.7)
 if body_style:s.top_margin=Cm(3.5);s.bottom_margin=Cm(3);s.left_margin=Cm(3.5);s.right_margin=Cm(2)
 s.header_distance=Cm(1.5);s.footer_distance=Cm(1.5)
 s.different_first_page_header_footer=False
 if body_style:
  borders=s._sectPr.find(qn('w:pgBorders'))
  if borders is not None:s._sectPr.remove(borders)
 for part in [s.header,s.footer,s.even_page_header,s.even_page_footer,s.first_page_header,s.first_page_footer]:
  part.is_linked_to_previous=False
  for p in part.paragraphs:p.clear()

configure_section(d.sections[0],False)
p=d.sections[0].footer.paragraphs[0];p.text='THÀNH PHỐ HỒ CHÍ MINH, NĂM 2026';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in p.runs:font(r,14,True)
configure_section(d.sections[1],True)
d.settings.odd_and_even_pages_header_footer=False

def para(txt='',style=None,align=None):
 p=d.add_paragraph(txt,style);p.paragraph_format.widow_control=True
 if align is not None:p.alignment=align
 if style is None:p.paragraph_format.first_line_indent=Cm(1)
 return p
def front_title(txt,break_before=True):
 p=para(txt,'Front Heading');p.paragraph_format.page_break_before=break_before;return p
def field(p,instr):
 r=p.add_run();b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');r._r.append(b)
 it=OxmlElement('w:instrText');it.set(qn('xml:space'),'preserve');it.text=' '+instr+' ';p.add_run()._r.append(it)
 se=OxmlElement('w:fldChar');se.set(qn('w:fldCharType'),'separate');p.add_run()._r.append(se)
 p.add_run('Đang cập nhật trong Word')
 en=OxmlElement('w:fldChar');en.set(qn('w:fldCharType'),'end');p.add_run()._r.append(en)

front_title('NHẬN XÉT CỦA GIẢNG VIÊN',False)
para('Học phần Ứng dụng AI',align=WD_ALIGN_PARAGRAPH.CENTER)
para('Sinh viên Mai Trần Thiện Tâm    MSSV 2305CT2070',align=WD_ALIGN_PARAGRAPH.CENTER)
para('Đề tài Xây dựng chatbot hỗ trợ tra cứu quy định pháp luật lao động Việt Nam')
para('Phần dành cho giảng viên nhận xét và đánh giá.')
for _ in range(10):para('........................................................................................................')
para('Điểm đánh giá .................................................................')
para('Giảng viên ký và ghi rõ họ tên',align=WD_ALIGN_PARAGRAPH.RIGHT)
front_title('LỜI CẢM ƠN')
para('Em xin cảm ơn quý thầy cô Trường Đại học Hùng Vương Thành phố Hồ Chí Minh đã giảng dạy các kiến thức nền tảng về lập trình, dữ liệu và trí tuệ nhân tạo. Những kiến thức này giúp em tiếp cận việc xây dựng chatbot hỗ trợ tra cứu quy định pháp luật lao động Việt Nam trong phạm vi học phần Ứng dụng AI.')
para('Em xin cảm ơn giảng viên phụ trách học phần [CẦN ĐIỀN HỌ TÊN] đã dành thời gian xem xét và góp ý cho báo cáo. Em mong tiếp tục nhận được nhận xét về phương pháp thực nghiệm, cách trình bày kết quả và những hạn chế cần khắc phục của sản phẩm.')
para('Báo cáo phân biệt các chức năng được xác định từ mã nguồn, kết quả kiểm thử đã chạy và những phần chưa có đủ minh chứng thực nghiệm. Em mong những giới hạn được trình bày rõ sẽ giúp việc đánh giá và hoàn thiện đề tài thuận lợi hơn.')
para('Mai Trần Thiện Tâm',align=WD_ALIGN_PARAGRAPH.RIGHT)
front_title('LỜI CAM ĐOAN')
para('Em là Mai Trần Thiện Tâm, MSSV 2305CT2070, trình bày báo cáo học phần Ứng dụng AI với đề tài “Xây dựng chatbot hỗ trợ tra cứu quy định pháp luật lao động Việt Nam”. Em chịu trách nhiệm kiểm tra nội dung báo cáo, xác định đúng phần việc mình thực hiện và ghi nhận các nguồn mã, dữ liệu, tài liệu hoặc công cụ đã sử dụng theo quy định của học phần.')
para('Các số liệu trong báo cáo được phân biệt giữa phép kiểm tra thực hiện ngày 12/10/2026 và kết quả kế thừa từ các tệp benchmark của dự án. Những phần chưa có nhật ký thực thi hoặc chưa được xác nhận, bao gồm fine-tuning và đánh giá đồ thị online, được ghi rõ giới hạn; không trình bày như kết quả đã nghiệm thu.')
para('Em cam kết trích dẫn nguồn tham khảo trung thực và không sử dụng các số liệu chưa được kiểm chứng để thay thế kết quả thực nghiệm. Những nội dung về quá trình thực hiện, quyền sử dụng mã nguồn và sự hỗ trợ nhận được cần được em đối chiếu trước khi ký bản nộp chính thức.')
para('TP. Hồ Chí Minh, ngày [CẦN ĐIỀN] tháng [CẦN ĐIỀN] năm 2026',align=WD_ALIGN_PARAGRAPH.RIGHT)
para('Người cam đoan',align=WD_ALIGN_PARAGRAPH.RIGHT)
para('\n\nMai Trần Thiện Tâm',align=WD_ALIGN_PARAGRAPH.RIGHT)
front_title('TÓM TẮT BÁO CÁO')
para('Báo cáo trình bày việc xây dựng VietLabor AI, một chatbot hỗ trợ tra cứu quy định pháp luật lao động Việt Nam. Hệ thống sử dụng Streamlit, mô hình Qwen2.5 qua Ollama và truy xuất lai giữa BM25 với embedding BGE-M3 lưu trong ChromaDB. Neo4j được tích hợp để mở rộng chứng cứ theo quan hệ giữa các điều luật khi dịch vụ đồ thị khả dụng. Pipeline có các bước phân rã vấn đề, kiểm tra tiền đề, lựa chọn chứng cứ, đóng gói ngữ cảnh và kiểm định trích dẫn.')
para('Tại thời điểm kiểm kê, corpus có 9.312 đoạn thuộc 36 mã văn bản. Kết quả ablation được lưu trong dự án ghi Hit@10 của hybrid là 95,86% và hybrid có chuẩn hóa là 97,93% trên 145 câu hỏi có thể chấm. Lần kiểm thử mới ghi nhận 366 passed, 6 skipped trong 46,40 giây. Báo cáo phân tích giới hạn của bộ 15 rubric, phân biệt khả năng kiểm soát trích dẫn với độ đúng pháp lý và chưa xác nhận kết quả fine-tuning hoặc lợi ích đồ thị online.')
para('Từ khóa: chatbot, pháp luật lao động, truy xuất lai, RAG, đồ thị tri thức, chứng cứ.')
front_title('ABSTRACT')
para('This report presents VietLabor AI, a chatbot for assisting users in looking up Vietnamese labour regulations. The system combines a Streamlit interface, Qwen2.5 inference through Ollama, BM25 lexical retrieval and BGE-M3 embeddings stored in ChromaDB. An optional Neo4j component expands evidence using relationships between statutory provisions. The pipeline includes issue decomposition, premise checking, evidence selection, context assembly and citation validation.')
para('The inspected local corpus contains 9,312 chunks associated with 36 document identifiers. A previously saved ablation reports a Hit@10 of 95.86% for hybrid retrieval and 97.93% for hybrid retrieval with query normalization on 145 scorable questions. A new default test run recorded 366 passed and 6 skipped tests in 46.40 seconds. The report discusses weaknesses in the existing rubric, distinguishes citation controls from legal correctness, and does not claim completed fine-tuning or measured benefits from an online knowledge graph.')
para('Keywords: chatbot, Vietnamese labour law, hybrid retrieval, RAG, knowledge graph, evidence.')
front_title('MỤC LỤC');field(para('','Field Container'),r'TOC \o "1-3" \h \z \u')
front_title('DANH MỤC CÁC KÝ HIỆU VÀ TỪ VIẾT TẮT')

table_count=0
def table(caption,headers,rows,widths=None):
 global table_count
 table_count+=1
 if caption:para(caption,'Bang')
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:
  if len(headers)==3:widths=[3.3,5.9,6.3]
  elif len(headers)==4:widths=[1.5,4.0,5.0,5.0]
  elif len(headers)==5:widths=[3.8,2.8,2.8,2.8,3.3]
  else:widths=[15.5/len(headers)]*len(headers)
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for c,txt in zip(cells,row):c.text=str(txt)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
  if ri==0:rep=OxmlElement('w:tblHeader');pr.append(rep)
  for ci,c in enumerate(row.cells):
   c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for edge,v in [('top','80'),('bottom','80'),('left','90'),('right','90')]:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),v);el.set(qn('w:type'),'dxa');mar.append(el)
   tcpr.append(mar)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcpr.append(borders)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6E6E6');tcpr.append(sh)
   for p in c.paragraphs:
    p.style=d.styles['Table Text'];p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if ri<3:p.paragraph_format.keep_with_next=True
    for run in p.runs:font(run,14,ri==0)
 spacer=para();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(1);spacer.paragraph_format.first_line_indent=Cm(0)
 font(spacer.add_run(''),1)
 return t

table('', ['Viết tắt','Tiếng Anh','Tiếng Việt'],[
 ['AI','Artificial Intelligence','Trí tuệ nhân tạo'],['LLM','Large Language Model','Mô hình ngôn ngữ lớn'],['RAG','Retrieval Augmented Generation','Sinh có tăng cường truy xuất'],['KG','Knowledge Graph','Đồ thị tri thức'],['RRF','Reciprocal Rank Fusion','Hợp nhất nghịch đảo thứ hạng'],['BM25','Best Matching 25','Thuật toán xếp hạng từ vựng'],['SFT','Supervised Fine Tuning','Tinh chỉnh có giám sát'],['LoRA','Low Rank Adaptation','Thích nghi hạng thấp'],['QLoRA','Quantized LoRA','LoRA với lượng tử hóa'],['OCR','Optical Character Recognition','Nhận dạng ký tự quang học'],['TTL','Time To Live','Thời gian sống của cache'],['MRR','Mean Reciprocal Rank','Trung bình nghịch đảo thứ hạng'],['UI','User Interface','Giao diện người dùng'],['HĐLĐ','Labour contract','Hợp đồng lao động'],['NLĐ','Employee','Người lao động'],['NSDLĐ','Employer','Người sử dụng lao động'],['BHXH','Social insurance','Bảo hiểm xã hội']],[2.2,6.4,6.9])
front_title('DANH MỤC CÁC HÌNH VẼ');field(para(),r'TOC \h \z \t "Hinh,1"')
front_title('DANH MỤC CÁC BẢNG BIỂU');field(para(),r'TOC \h \z \t "Bang,1"')

main=d.add_section(WD_SECTION_START.NEW_PAGE);configure_section(main,True)
pg=main._sectPr.find(qn('w:pgNumType'))
if pg is None:pg=OxmlElement('w:pgNumType');main._sectPr.append(pg)
pg.set(qn('w:start'),'1');pg.set(qn('w:fmt'),'decimal')
p=main.header.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(p,'PAGE')
for r in p.runs:font(r,12)

def figure(caption,filename):
 p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
 path=ROOT/'ui/static/approved-homepage.png' if filename=='approved-homepage.png' else OUT/('minh_chung' if filename.startswith('UI') else 'so_do')/filename
 from PIL import Image
 with Image.open(path) as im:width,height=im.size
 target_w=15.5;target_h=target_w*height/width
 max_h=13 if filename.startswith('UI') else 16
 if target_h>max_h:target_w*=max_h/target_h
 p.add_run().add_picture(str(path),width=Cm(target_w))
 p=para(caption,'Hinh');p.paragraph_format.keep_with_next=True
 src='Nguồn ảnh tham chiếu có sẵn ui/static/approved-homepage.png; không xác nhận lần chạy thực tế.' if filename=='approved-homepage.png' else ('Nguồn ảnh chụp ứng dụng tại 127.0.0.1:8501, ngày 12/10/2026; giữ nguyên giao diện quan sát.' if filename.startswith('UI') else 'Nguồn tác giả tổng hợp từ mã nguồn dự án và số liệu được ghi trong nội dung.')
 p=para(src);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.0
 for r in p.runs:font(r,11,False,True)

lines=(OUT/'noi_dung_bao_cao.md').read_text(encoding='utf-8').splitlines();i=0;first_heading=True
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line.startswith('@table '):
  caption='Bảng '+line[7:];raw=[]
  while i<len(lines) and lines[i].strip()!='@endtable':
   raw.append([x.strip() for x in lines[i].split('|')]);i+=1
  i+=1
  widths=None
  if caption.startswith('Bảng 3.1'):widths=[1.5,6.1,7.9]
  table(caption,raw[0],raw[1:],widths)
 elif line.startswith('@figure '):
  cap,name=line[8:].split(' | ');figure('Hình '+cap,name)
 elif line=='@code':
  while i<len(lines) and lines[i].strip()!='@endcode':
   p=para(lines[i],'Code');p.paragraph_format.first_line_indent=Cm(0);i+=1
  i+=1
 elif line=='@corpus_table':
  rows=[[str(j+1),x['doc_id'],x['number'] or '',x['count']] for j,x in enumerate(audit['corpus']['docs'])]
  table('Bảng A.1 Danh mục văn bản trong corpus',['TT','Mã văn bản','Số hiệu','Đoạn'],rows,[1,5.6,7.1,1.8])
 elif line=='@training_table':
  table('Bảng A.2 Số bản ghi trong các tệp ChatML',['Tệp','Bản ghi','Ghi chú'],[[k,v['count'],'Chưa kiểm tra trùng giữa các tệp'] for k,v in audit['training'].items()],[8.8,1.6,5.1])
 elif line=='@corpus_hash':
  p=para(audit['corpus']['sha256'],'Code');p.paragraph_format.first_line_indent=Cm(0)
 elif line=='@rubric_table':
  raw=(ROOT/'reports/rubric_15_test_cases.md').read_text(encoding='utf-8').splitlines()
  rows=[]
  for row in raw:
   parts=[x.strip().replace('`','') for x in row.strip('|').split('|')]
   if len(parts)==8 and parts[0].isdigit():rows.append([parts[1],parts[3],parts[4]])
  table('Bảng B.1 Danh sách câu hỏi rubric và mục tiêu chấm',['Mã ca','Đầu vào','Mục tiêu chấm'],rows,[2.5,6.5,6.5])
 elif line.startswith('# '):
  p=para(line[2:],'Heading 1')
  for r in p.runs:font(r,16,True)
  if first_heading:p.paragraph_format.page_break_before=False;first_heading=False
 elif line.startswith('## '):
  p=para(line[3:].upper(),'Heading 2')
  if re.match(r'## F\.[2-5] ',line):p.paragraph_format.page_break_before=True
  for r in p.runs:font(r,14,True)
 elif line.startswith('### '):
  p=para(line[4:],'Heading 3')
  for r in p.runs:font(r,14,True)
 else:para(line)

settings=d.settings._element;up=settings.find(qn('w:updateFields'))
if up is None:up=OxmlElement('w:updateFields');settings.append(up)
up.set(qn('w:val'),'true')
d.core_properties.title='Xây dựng chatbot hỗ trợ tra cứu quy định pháp luật lao động Việt Nam'
d.core_properties.author='Mai Trần Thiện Tâm';d.core_properties.subject='Báo cáo học phần Ứng dụng AI';d.core_properties.comments='Bản soạn theo mã nguồn ngày 12/10/2026. Cần bổ sung thông tin học vụ và các minh chứng được ghi trong phụ lục.'
d.save(FINAL)
print(FINAL)
print('Sections',len(d.sections),'Tables',len(d.tables),'Pictures',len(d.inline_shapes),'Paragraphs',len(d.paragraphs))
