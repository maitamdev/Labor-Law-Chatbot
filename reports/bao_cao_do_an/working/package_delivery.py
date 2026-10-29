from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import pymupdf,json,hashlib,zipfile,shutil,sys
sys.stdout.reconfigure(encoding='utf-8')
root=Path(r'D:\chatbot-law');out=root/'reports/bao_cao_do_an';w=out/'working';ev=out/'minh_chung'
report=out/'Bao_cao_Hoc_phan_Ung_dung_AI_Mai_Tran_Thien_Tam.docx'
audit=json.loads((w/'audit.json').read_text(encoding='utf-8'))
qa=json.loads((w/'page_audit.json').read_text(encoding='utf-8'))
d=Document(report);pdf=pymupdf.open(w/'bao_cao_qa.pdf')
def exact_page(label):
 return next(i+1 for i,p in enumerate(pdf) if label in [x.strip() for x in p.get_text().splitlines()])
start=exact_page('MỞ ĐẦU');refs=exact_page('TÀI LIỆU THAM KHẢO')
captions=[p.text for p in d.paragraphs if p.style.name in ['Hinh','Bang']]
chapters=[p.text for p in d.paragraphs if p.style.name=='Heading 1' and p.text.startswith('CHƯƠNG ')]
assert len(chapters)==6
assert len([c for c in captions if c.startswith('Hình ')])==15
assert len([c for c in captions if c.startswith('Bảng ')])==19
assert not any(not x['chars'] for x in qa['audit']), 'Blank page'
assert not any(x['outside'] for x in qa['audit']), 'Text outside page limits'
assert len(pdf)==103 and refs-start==66
assert len(d.inline_shapes)==17
fields=[e.text or '' for e in d._element.iter(qn('w:instrText'))]
assert sum('TOC ' in f for f in fields)==3
for s in list(d.sections)[1:]:
 assert abs(s.left_margin.cm-3.5)<.01 and abs(s.right_margin.cm-2)<.01
 assert abs(s.top_margin.cm-3.5)<.01 and abs(s.bottom_margin.cm-3)<.01
source=Path(r'C:\Users\Asus\Desktop\Bao_cao_Do_an_co_so_POS_Quan_ly_kho_Canh_bao_ton_kho_thap_FINAL.docx')
assert hashlib.sha256(source.read_bytes()).hexdigest()==audit['pos']['sha256']
manifest={
 'date':'2026-10-12','student':'Mai Trần Thiện Tâm','student_id':'2305CT2070',
 'report_sha256':hashlib.sha256(report.read_bytes()).hexdigest(),
 'pages_total':len(pdf),'front_pages':start-1,'main_text_pages_including_figures_tables':refs-start,
 'references_pages':2,'appendix_pages':len(pdf)-(refs+1),'main_text_requirement_70_to_100_met':False,
 'chapters':6,'appendices':6,'technical_diagrams':7,'charts':2,'runtime_screenshots':6,
 'figure_captions_in_report':15,'table_captions':19,'toc_fields':3,
 'pytest':{'passed':366,'skipped':6,'failed':0,'seconds':46.40,'command':'.venv\\Scripts\\python.exe -m pytest tests -q'},
 'ui_demo':{'local_model':'qwen2.5:7b-instruct-q4_0','graph_online':False,
 'in_scope':'Returned 30 days with Article 35 clause 1 point b visible',
 'clarification':'Requested contract type',
 'followup':'Generic processing failure after selecting 12-36 month contract option; unresolved',
 'out_of_scope':'Rejected land inheritance question'},
 'qa':{'renderer':'Microsoft Word COM + PyMuPDF','reason':'Bundled render_docx.py unavailable: missing soffice.exe',
 'all_pages_rendered_and_visually_reviewed':True,'blank_pages':[],'text_outside_page_limits':[]},
 'original_pos_docx_unchanged':True,
 'remaining':['lecturer/class/cohort/signature date','followup error','temporal metadata and rubric review','file attachment demo','online graph experiment','SFT execution evidence if required','70-100 text-page requirement'],
}
(ev/'kiem_tra_ban_giao.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
shutil.copy2(w/'pytest_run.txt',ev/'pytest_run.txt')
shutil.copy2(w/'audit.json',ev/'kiem_ke_du_lieu.json')
readme=f'''# Hồ sơ bàn giao báo cáo học phần Ứng dụng AI

Mai Trần Thiện Tâm — 2305CT2070. Đề tài: Xây dựng chatbot hỗ trợ tra cứu quy định pháp luật lao động Việt Nam.

Báo cáo Word có {len(pdf)} trang sau khi dàn bằng Microsoft Word: {start-1} trang đầu, {refs-start} trang từ Mở đầu đến hết Chương 6 (có hình/bảng), 2 trang tham khảo và {len(pdf)-(refs+1)} trang phụ lục. Chưa đạt yêu cầu 70–100 trang văn bản loại trừ hình/bảng/tham khảo nếu áp dụng nguyên văn biểu mẫu.

- Bao_cao_Hoc_phan_Ung_dung_AI_Mai_Tran_Thien_Tam.docx: bản báo cáo để đọc/chỉnh sửa; còn [CẦN ĐIỀN].
- DOI_CHIEU_YEU_CAU_VA_VIEC_CAN_LAM.md: phần đã làm, các yêu cầu và việc cần hoàn thiện.
- so_do/: 7 sơ đồ và 2 biểu đồ, có PNG/SVG/PDF; bảy sơ đồ có Mermaid .mmd.
- minh_chung/: 6 ảnh runtime, 4 tệp đầu ra của ca demo, kết quả pytest và kiểm kê dữ liệu. Demo có một lỗi nối tiếp chưa được sửa.
- nguon_benchmark/: các bảng/CSV đã có trước, không phải số liệu chạy mới trong lần soạn báo cáo.
- noi_dung_bao_cao.md và ma_tao_tai_lieu/: nguồn nội dung và mã tạo; chạy lại yêu cầu workspace gốc, báo cáo POS tham chiếu, python-docx, Pillow, ReportLab, PyMuPDF và Microsoft Word trên Windows. add_demo_content.py là thao tác biên soạn đã áp dụng, không chạy lặp.

Các bản PDF của sơ đồ dùng để in hình riêng. Không kèm PDF báo cáo và PNG từng trang phục vụ QA, corpus đầy đủ, model, thông tin kết nối hoặc lịch sử hội thoại cũ.

Kết quả pytest: 366 passed, 6 skipped in 46.40s. Notebook SFT chưa có output xác nhận; Neo4j chưa online trong demo; không công bố độ chính xác pháp lý 100%.
'''
(out/'HUONG_DAN_SU_DUNG.md').write_text(readme,encoding='utf-8')
zip_path=out/'Ho_so_bao_cao_Ung_dung_AI.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED,compresslevel=7) as z:
 for name in [report.name,'DOI_CHIEU_YEU_CAU_VA_VIEC_CAN_LAM.md','HUONG_DAN_SU_DUNG.md','noi_dung_bao_cao.md']:
  z.write(out/name,name)
 for f in (out/'so_do').iterdir():
  if f.suffix in ['.png','.svg','.pdf'] or (f.suffix=='.mmd' and int(f.name[:2])<=7):z.write(f,'so_do/'+f.name)
 for f in ev.iterdir():
  if f.suffix in ['.jpg','.txt','.json']:z.write(f,'minh_chung/'+f.name)
 for name in ['ablation_retrieval.md','ablation_retrieval.csv','benchmark_11_gold_results.md','rubric_15_test_cases.md']:
  z.write(root/'reports'/name,'nguon_benchmark/'+name)
 for name in ['build_report.py','build_diagrams.py','render_word.ps1','artifact.md','qa_render.py']:
  z.write(w/name,'ma_tao_tai_lieu/'+name)
with zipfile.ZipFile(zip_path) as z:
 assert z.testzip() is None
 assert z.read(report.name)==report.read_bytes()
 print('ZIP verified:',len(z.namelist()),'files,',zip_path.stat().st_size,'bytes')
print(json.dumps(manifest,ensure_ascii=False,indent=2))
