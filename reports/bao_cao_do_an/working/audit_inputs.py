from pathlib import Path
import json, hashlib, ast, collections
from docx import Document
import fitz

root=Path(r'D:\chatbot-law'); out=root/'reports/bao_cao_do_an/working'
summary={}
for label,path in [('template',out/'bieu_mau.docx'),('pos',Path(r'C:\Users\Asus\Desktop\Bao_cao_Do_an_co_so_POS_Quan_ly_kho_Canh_bao_ton_kho_thap_FINAL.docx'))]:
 d=Document(path)
 texts=[p.text for p in d.paragraphs]
 texts += [c.text for t in d.tables for row in t.rows for c in row.cells]
 (out/(label+'_extracted.txt')).write_text('\n'.join(texts),encoding='utf-8')
 summary[label]={'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sections':[{'width_cm':s.page_width.cm,'height_cm':s.page_height.cm,'left_cm':s.left_margin.cm,'right_cm':s.right_margin.cm,'top_cm':s.top_margin.cm,'bottom_cm':s.bottom_margin.cm} for s in d.sections], 'paragraphs':len(d.paragraphs),'tables':len(d.tables),'images':len(d.inline_shapes),'headings':[p.text for p in d.paragraphs if p.style and p.style.name.startswith('Heading')]}
 pdf=fitz.open(out/('bieu_mau.pdf' if label=='template' else 'pos_reference.pdf'))
 summary[label]['pages']=len(pdf)
 folder=out/(label+'_pages');folder.mkdir(exist_ok=True)
 for i,p in enumerate(pdf):
  p.get_pixmap(matrix=fitz.Matrix(1,1)).save(folder/f'page-{i+1}.png')
corpus=root/'data/processed/legal_documents_v3.jsonl'
records=[json.loads(l) for l in corpus.read_text(encoding='utf-8').splitlines() if l.strip()]
docs={}
for c in records:
 k=c['doc_id']
 if k not in docs: docs[k]={'doc_id':k, 'title':c.get('doc_title'),'number':c.get('document_no'),'status':c.get('status'),'effective_from':c.get('effective_from'),'effective_to':c.get('effective_to'),'source':c.get('official_source'),'count':0}
 docs[k]['count']+=1
summary['corpus']={'chunks':len(records),'documents':len(docs),'unique_ids':len(set(c['chunk_id'] for c in records)), 'sha256':hashlib.sha256(corpus.read_bytes()).hexdigest(),'statuses':dict(collections.Counter(c.get('status') for c in records)),'domains':dict(collections.Counter(c.get('domain') for c in records)),'docs':list(docs.values())}
summary['training']={}
for p in (root/'data/training').glob('*'):
 if p.suffix=='.jsonl':
  rows=[json.loads(l) for l in p.read_text(encoding='utf-8').splitlines() if l.strip()]
  summary['training'][p.name]={'count':len(rows),'keys':list(rows[0]) if rows else [],'roles':list(set(m.get('role','') for r in rows for m in r.get('messages',[])))}
nb=json.loads((root/'VietLabor_SFT_Qwen25_Colab.ipynb').read_text(encoding='utf-8'))
code='\n\n'.join(''.join(c['source']) for c in nb['cells'] if c['cell_type']=='code')
(out/'notebook_code.txt').write_text(code,encoding='utf-8')
summary['notebook']={'code_cells':sum(c['cell_type']=='code' for c in nb['cells']),'executed_cells':sum(c.get('execution_count') is not None for c in nb['cells'] if c['cell_type']=='code'),'output_blocks':sum(len(c.get('outputs',[])) for c in nb['cells'])}
schema=ast.parse((root/'graph/schema.py').read_text(encoding='utf-8'))
for node in ast.walk(schema):
 if isinstance(node,ast.AnnAssign) and isinstance(node.target,ast.Name) and node.target.id=='DEFAULT_STATUTORY_BRIDGES': summary['bridges_defined']=len(node.value.elts)
summary['source_files']={}
for folder in ['app','rag','graph','ingestion','models','ui','evaluation','tests']:
 summary['source_files'][folder]=len(list((root/folder).rglob('*.py')))
(out/'audit.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in summary.items() if k not in ['pos','template','corpus']},ensure_ascii=False,indent=2))
print('Corpus',len(records),'chunks',len(docs),'documents')
print('Reference pages',summary['template']['pages'],summary['pos']['pages'])
